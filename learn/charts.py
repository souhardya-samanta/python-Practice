import openpyxl as xl
from openpyxl.chart import BarChart, Reference

wb = xl.load_workbook("head.xlsx")
sheet = wb['Sheet1']

values = Reference(sheet, min_row=2, max_row=sheet.max_row, min_col=5, max_col=5)

chart = BarChart()
chart.add_data(values)
sheet.add_chart(chart, "i2")
wb.save("chart.xlsx")
